# -*- coding: utf-8 -*-
"""Word 排版工具：A4 版心、中文字体、表格、页脚。"""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x8C, 0x1F, 0x28)
INK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x59, 0x59, 0x59)
GOLD = RGBColor(0xB0, 0x86, 0x3A)

HEI, SONG = "微软雅黑", "宋体"


def _font(run, name, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), name)
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def new_doc(margins=(2.0, 2.0, 2.2, 2.2), footer_text=None, body=10.5):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(margins[0]), Cm(margins[1])
    sec.left_margin, sec.right_margin = Cm(margins[2]), Cm(margins[3])

    st = doc.styles["Normal"]
    st.font.name = SONG
    st.font.size = Pt(body)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
    st.paragraph_format.line_spacing = 1.34
    st.paragraph_format.space_after = Pt(4)

    if footer_text:
        para = sec.footer.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = para.add_run(footer_text)
        _font(r, HEI, 8, color=RGBColor(0x99, 0x99, 0x99))
    return doc


def title(doc, text, sub=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    _font(r, HEI, 20, True, ACCENT)
    if sub:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(10)
        r2 = p2.add_run(sub)
        _font(r2, HEI, 9.5, color=GREY)
    # 分隔线
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), "8C1F28")
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def h1(doc, text, num=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    if num:
        r0 = p.add_run(num + "  ")
        _font(r0, HEI, 12, True, GOLD)
    r = p.add_run(text)
    _font(r, HEI, 12, True, ACCENT)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    _font(r, HEI, 10.5, True, INK)
    return p


def p(doc, text, indent=False, size=10.5, color=INK, bold=False, space_after=5, align=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    if indent:
        para.paragraph_format.first_line_indent = Pt(size * 2)
    if align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _rich(para, text, size, color, bold)
    return para


def _rich(para, text, size, color, bold):
    """支持 **加粗** 与 `等宽` 两种行内标记。"""
    import re
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|`(.+?)`", text):
        if m.start() > pos:
            r = para.add_run(text[pos:m.start()])
            _font(r, SONG, size, bold, color)
        if m.group(1) is not None:
            r = para.add_run(m.group(1))
            _font(r, HEI, size, True, ACCENT)
        else:
            r = para.add_run(m.group(2))
            _font(r, "Consolas", size - 0.5, False, RGBColor(0x44, 0x44, 0x44))
        pos = m.end()
    if pos < len(text):
        r = para.add_run(text[pos:])
        _font(r, SONG, size, bold, color)


def bullet(doc, text, size=10.5, marker="·"):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(0.75)
    para.paragraph_format.first_line_indent = Cm(-0.35)
    para.paragraph_format.space_after = Pt(3)
    r = para.add_run(marker + " ")
    _font(r, HEI, size, True, GOLD)
    _rich(para, text, size, INK, False)
    return para


def num_item(doc, idx, text, size=10.5):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(0.95)
    para.paragraph_format.first_line_indent = Cm(-0.55)
    para.paragraph_format.space_after = Pt(4)
    r = para.add_run("%d）" % idx)
    _font(r, HEI, size, True, ACCENT)
    _rich(para, text, size, INK, False)
    return para


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def _set_cell(cell, text, size=9.5, bold=False, color=INK, align=None, font=SONG):
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.line_spacing = 1.15
    if align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _rich(para, text, size, color, bold)
    for r in para.runs:
        if r.font.name == SONG:
            _font(r, font, size, r.font.bold, color)


def table(doc, headers, rows, widths, size=9.5, header_fill="8C1F28", zebra=True, align_center=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    align_center = align_center or []
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = Cm(widths[i])
        _set_cell(c, h, size, True, RGBColor(0xFF, 0xFF, 0xFF), "center", HEI)
        shade(c, header_fill)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].width = Cm(widths[i])
            _set_cell(cells[i], str(v), size, False, INK, "center" if i in align_center else None)
            if zebra and ri % 2 == 1:
                shade(cells[i], "F7F1EC")
    return t


def quote(doc, text, size=10):
    """左侧色条引用块，用于「可直接复制的提示词」这类内容。"""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    cell.width = Cm(16.6)
    cell.text = ""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), "B0863A")
    for edge in ("top", "bottom", "right"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "nil")
        borders.append(e)
    borders.append(left)
    tcPr.append(borders)
    shade(cell, "FAF5EF")
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(4)
    _rich(para, text, size, RGBColor(0x33, 0x33, 0x33), False)
    return t


def spacer(doc, pt=6):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1.0
    r = para.add_run("")
    _font(r, SONG, pt)
